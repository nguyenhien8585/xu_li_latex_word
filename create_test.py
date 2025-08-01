#!/usr/bin/env python3
"""
📝 Script tạo file đề thi mẫu để test ứng dụng
Phiên bản đơn giản
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_simple_exam():
    """Tạo file đề thi mẫu đơn giản"""
    
    doc = Document()
    
    # Tiêu đề
    title = doc.add_heading('ĐỀ THI MẪU - TOÁN HỌC', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Thông tin đề thi
    info = doc.add_paragraph('Thời gian: 90 phút | Số câu: 10 câu')
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # Khoảng trống
    
    # PHẦN I: TRẮC NGHIỆM
    doc.add_heading('PHẦN I: CÂU HỎI TRẮC NGHIỆM (5 câu)', level=2)
    
    # Câu trắc nghiệm mẫu
    questions_text = """
Câu 1. Tập xác định của hàm số y = √(x-1) là:
A. [1; +∞)
B. (-∞; 1]
C. (1; +∞)
D. ℝ

Câu 2. Giá trị của lim(x→∞) (2x+1)/(x-3) là:
A. 0
B. 1
C. 2
D. +∞

Câu 3. Đạo hàm của hàm số f(x) = x³ - 2x + 1 là:
A. 3x² - 2
B. 3x² + 2
C. x² - 2x
D. 3x - 2

Câu 4. Phương trình x² - 5x + 6 = 0 có nghiệm là:
A. x = 1; x = 6
B. x = 2; x = 3
C. x = -2; x = -3
D. Vô nghiệm

Câu 5. Giá trị nhỏ nhất của hàm số y = x² - 4x + 5 trên ℝ là:
A. 1
B. 2
C. 3
D. 5
"""
    
    doc.add_paragraph(questions_text)
    
    # PHẦN II: ĐÚNG/SAI
    doc.add_heading('PHẦN II: CÂU HỎI ĐÚNG/SAI (3 câu)', level=2)
    
    tf_text = """
Câu 6. Hàm số y = sin(x) có chu kỳ là 2π.

Câu 7. Đạo hàm của hàm số y = e^x là e^x.

Câu 8. Tích phân ∫₀¹ x dx = 1/2.
"""
    
    doc.add_paragraph(tf_text)
    
    # PHẦN III: TỰ LUẬN
    doc.add_heading('PHẦN III: CÂU HỎI TỰ LUẬN (2 câu)', level=2)
    
    essay_text = """
Câu 9. Cho hàm số y = x³ - 3x² + 2. Tìm cực trị của hàm số và vẽ đồ thị.

Câu 10. Giải phương trình lượng giác: 2sin²x + 3cosx - 3 = 0 trên [0; 2π].
"""
    
    doc.add_paragraph(essay_text)
    
    # Lưu file
    doc.save('test_exam.docx')
    print("✅ Đã tạo file 'test_exam.docx'")
    print("🚀 Có thể sử dụng file này để test ứng dụng!")

def create_math_exam():
    """Tạo đề thi có công thức LaTeX đơn giản"""
    
    doc = Document()
    
    title = doc.add_heading('ĐỀ THI TOÁN - CÓ CÔNG THỨC', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Câu hỏi với LaTeX
    math_text = """
Câu 1. Tính giới hạn $\\lim_{x \\to 0} \\frac{\\sin x}{x}$:
A. 0
B. 1
C. $+\\infty$
D. Không tồn tại

Câu 2. Đạo hàm của $f(x) = x^2 + 3x - 1$ là:
A. $f'(x) = 2x + 3$
B. $f'(x) = x^2 + 3$
C. $f'(x) = 2x - 1$
D. $f'(x) = x + 3$

Câu 3. Công thức tính diện tích hình tròn bán kính r:
A. $S = \\pi r$
B. $S = \\pi r^2$
C. $S = 2\\pi r$
D. $S = \\frac{\\pi r^2}{2}$
"""
    
    doc.add_paragraph(math_text)
    
    # Bảng Markdown
    table_text = """
Câu 4. Cho bảng thống kê điểm số:

| Điểm | Số HS | Tỷ lệ |
|------|-------|-------|
| 0-5  | 10    | 20%   |
| 6-7  | 25    | 50%   |
| 8-10 | 15    | 30%   |

Điểm trung bình gần nhất với:
A. 6.5
B. 7.0
C. 7.5
D. 8.0
"""
    
    doc.add_paragraph(table_text)
    
    doc.save('math_exam_latex.docx')
    print("✅ Đã tạo file 'math_exam_latex.docx'")
    print("🧮 File này có chứa công thức LaTeX và bảng Markdown!")

if __name__ == "__main__":
    print("🚀 Tạo file đề thi mẫu...")
    print()
    
    # Tạo file đơn giản
    create_simple_exam()
    print()
    
    # Tạo file có LaTeX
    create_math_exam()
    print()
    
    print("✨ Hoàn thành! Có thể upload các file .docx vào ứng dụng để test.")
    print("📁 Files được tạo:")
    print("   - test_exam.docx (đề thi cơ bản)")
    print("   - math_exam_latex.docx (có LaTeX + bảng)")
    print()
    print("🔥 Chạy ứng dụng: streamlit run app.py")
