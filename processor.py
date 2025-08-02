#!/usr/bin/env python3
"""
🧮 Module xử lý LaTeX và Markdown đơn giản và ổn định
Tạo file Word có thể mở được 100%
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from datetime import datetime

def detect_latex_patterns(text):
    """Phát hiện các pattern LaTeX trong text"""
    
    patterns = [
        (r'\$\{([^}]+)\}\$', 'inline'),      # ${...}$
        (r'\$([^$]+)\$', 'inline'),          # $...$
        (r'\\?\\\(([^)]+)\\\)', 'inline'),   # \(...\)
        (r'\\?\\\[([^\]]+)\\\]', 'display')  # \[...\]
    ]
    
    found_formulas = []
    
    for pattern, math_type in patterns:
        matches = re.finditer(pattern, text)
        for match in matches:
            latex_content = match.group(1)
            found_formulas.append({
                'original': match.group(0),
                'latex': latex_content,
                'type': math_type,
                'start': match.start(),
                'end': match.end()
            })
    
    # Sắp xếp theo vị trí
    found_formulas.sort(key=lambda x: x['start'], reverse=True)
    
    return found_formulas

def detect_markdown_tables(text):
    """Phát hiện bảng Markdown trong text"""
    
    # Pattern đơn giản hơn cho bảng Markdown
    lines = text.split('\n')
    tables = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Kiểm tra có phải dòng header không
        if '|' in line and line.count('|') >= 2:
            # Kiểm tra dòng tiếp theo có phải separator không
            if i + 1 < len(lines) and '---' in lines[i + 1] and '|' in lines[i + 1]:
                # Tìm các dòng data
                header = [cell.strip() for cell in line.strip('|').split('|')]
                rows = []
                
                j = i + 2  # Bỏ qua dòng separator
                while j < len(lines):
                    data_line = lines[j].strip()
                    if data_line and '|' in data_line:
                        cells = [cell.strip() for cell in data_line.strip('|').split('|')]
                        # Đảm bảo số cột đúng
                        while len(cells) < len(header):
                            cells.append('')
                        rows.append(cells[:len(header)])
                        j += 1
                    else:
                        break
                
                if header and rows:
                    tables.append({
                        'header': header,
                        'rows': rows,
                        'start_line': i,
                        'end_line': j
                    })
                
                i = j
            else:
                i += 1
        else:
            i += 1
    
    return tables

def convert_latex_to_readable(latex_code):
    """Chuyển LaTeX thành text dễ đọc"""
    
    # Các chuyển đổi cơ bản
    conversions = {
        r'\\frac\{([^}]+)\}\{([^}]+)\}': r'(\1)/(\2)',  # Fraction
        r'\^(\w)': r'^(\1)',  # Superscript đơn
        r'\^{([^}]+)}': r'^(\1)',  # Superscript phức tạp
        r'_(\w)': r'_\1',  # Subscript đơn  
        r'_{([^}]+)}': r'_(\1)',  # Subscript phức tạp
        r'\\sqrt\{([^}]+)\}': r'√(\1)',  # Square root
        r'\\sum': '∑',  # Summation
        r'\\int': '∫',  # Integral
        r'\\infty': '∞',  # Infinity
        r'\\alpha': 'α', r'\\beta': 'β', r'\\gamma': 'γ',
        r'\\delta': 'δ', r'\\pi': 'π', r'\\theta': 'θ',
        r'\\lambda': 'λ', r'\\mu': 'μ', r'\\sigma': 'σ',
        r'\\leq': '≤', r'\\geq': '≥', r'\\neq': '≠',
        r'\\pm': '±', r'\\times': '×', r'\\div': '÷',
        r'\\rightarrow': '→', r'\\leftarrow': '←',
        r'\\Rightarrow': '⇒', r'\\Leftarrow': '⇐'
    }
    
    result = latex_code
    for pattern, replacement in conversions.items():
        result = re.sub(pattern, replacement, result)
    
    return result

def process_paragraph_with_latex(doc, text):
    """Xử lý đoạn văn có LaTeX và tạo paragraph mới"""
    
    formulas = detect_latex_patterns(text)
    
    if not formulas:
        # Không có công thức, tạo paragraph thường
        return doc.add_paragraph(text)
    
    # Có công thức, xử lý từng phần
    para = doc.add_paragraph()
    current_pos = 0
    
    # Sắp xếp lại theo thứ tự xuất hiện
    formulas.sort(key=lambda x: x['start'])
    
    for formula in formulas:
        # Thêm text trước công thức
        if formula['start'] > current_pos:
            before_text = text[current_pos:formula['start']]
            if before_text:
                para.add_run(before_text)
        
        # Thêm công thức (dạng text được format)
        readable_formula = convert_latex_to_readable(formula['latex'])
        formula_run = para.add_run(f"[{readable_formula}]")
        formula_run.font.italic = True
        formula_run.font.bold = True
        formula_run.font.color.rgb = RGBColor(0, 100, 139)  # Màu xanh đậm
        formula_run.font.size = Pt(11)
        
        current_pos = formula['end']
    
    # Thêm text cuối
    if current_pos < len(text):
        remaining_text = text[current_pos:]
        if remaining_text:
            para.add_run(remaining_text)
    
    return para

def create_simple_table(doc, header, rows):
    """Tạo bảng Word đơn giản và ổn định"""
    
    if not header or not rows:
        return None
    
    # Tạo bảng
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Thiết lập header
    header_cells = table.rows[0].cells
    for i, header_text in enumerate(header):
        cell = header_cells[i]
        cell.text = header_text
        
        # Định dạng header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(11)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Màu nền header (đơn giản)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        try:
            shading_elm = parse_xml(r'<w:shd {} w:fill="366092"/>'.format(
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'))
            cell._tc.get_or_add_tcPr().append(shading_elm)
        except:
            pass  # Nếu không thể thêm màu, bỏ qua
    
    # Thêm dữ liệu
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            if i < len(row_cells):
                cell = row_cells[i]
                
                # Kiểm tra nếu cell có LaTeX
                if '$' in cell_data or '\\(' in cell_data:
                    # Xử lý LaTeX trong cell đơn giản
                    formulas = detect_latex_patterns(cell_data)
                    if formulas:
                        # Thay thế LaTeX bằng text dễ đọc
                        processed_text = cell_data
                        for formula in sorted(formulas, key=lambda x: x['start'], reverse=True):
                            readable = convert_latex_to_readable(formula['latex'])
                            processed_text = processed_text[:formula['start']] + f"[{readable}]" + processed_text[formula['end']:]
                        
                        cell.text = processed_text
                        
                        # Format text trong cell
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if '[' in run.text and ']' in run.text:
                                    run.font.italic = True
                                    run.font.color.rgb = RGBColor(0, 100, 139)
                    else:
                        cell.text = cell_data
                else:
                    cell.text = cell_data
                
                # Định dạng cell
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    
    return table

def process_content_simple(doc_content):
    """Xử lý nội dung document đơn giản"""
    
    # Tạo document mới
    new_doc = Document()
    
    # Thiết lập margins
    sections = new_doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Header
    header = new_doc.add_heading('DOCUMENT ĐÃ CHUYỂN ĐỔI', level=1)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.color.rgb = RGBColor(31, 78, 121)
        run.font.size = Pt(16)
    
    # Thông tin
    info_para = new_doc.add_paragraph(f"Ngày chuyển đổi: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in info_para.runs:
        run.font.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(128, 128, 128)
    
    note_para = new_doc.add_paragraph("📝 LaTeX được chuyển thành text có format, bảng Markdown thành bảng Word")
    note_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in note_para.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
    
    new_doc.add_paragraph()
    
    # Thống kê
    stats = {
        'latex_count': 0,
        'table_count': 0,
        'processed_paragraphs': 0
    }
    
    # Xử lý từng paragraph
    for para in doc_content.paragraphs:
        text = para.text.strip()
        
        if not text:
            new_doc.add_paragraph()  # Giữ khoảng trống
            continue
        
        stats['processed_paragraphs'] += 1
        
        # Kiểm tra có bảng Markdown không
        tables = detect_markdown_tables(text)
        if tables:
            stats['table_count'] += len(tables)
            
            # Xử lý text và bảng
            lines = text.split('\n')
            current_line = 0
            
            for table in tables:
                # Text trước bảng
                if table['start_line'] > current_line:
                    before_lines = lines[current_line:table['start_line']]
                    before_text = '\n'.join(before_lines).strip()
                    if before_text:
                        process_paragraph_with_latex(new_doc, before_text)
                
                # Tạo bảng
                table_obj = create_simple_table(new_doc, table['header'], table['rows'])
                if table_obj:
                    new_doc.add_paragraph()  # Khoảng cách sau bảng
                
                current_line = table['end_line']
            
            # Text sau bảng cuối
            if current_line < len(lines):
                remaining_lines = lines[current_line:]
                remaining_text = '\n'.join(remaining_lines).strip()
                if remaining_text:
                    process_paragraph_with_latex(new_doc, remaining_text)
        else:
            # Không có bảng, chỉ xử lý LaTeX
            formulas = detect_latex_patterns(text)
            if formulas:
                stats['latex_count'] += len(formulas)
            
            process_paragraph_with_latex(new_doc, text)
    
    # Thêm thống kê cuối document
    new_doc.add_paragraph()
    stats_para = new_doc.add_paragraph("📊 THỐNG KÊ CHUYỂN ĐỔI:")
    stats_para.runs[0].font.bold = True
    stats_para.runs[0].font.size = Pt(12)
    
    new_doc.add_paragraph(f"• Đã xử lý: {stats['processed_paragraphs']} đoạn văn")
    new_doc.add_paragraph(f"• Công thức LaTeX: {stats['latex_count']} công thức")
    new_doc.add_paragraph(f"• Bảng Markdown: {stats['table_count']} bảng")
    
    return new_doc, stats

def convert_docx_simple(uploaded_file):
    """Chuyển đổi file Word đơn giản và ổn định"""
    
    try:
        # Đọc document gốc
        doc = Document(uploaded_file)
        
        # Xử lý nội dung
        new_doc, stats = process_content_simple(doc)
        
        return new_doc, stats
        
    except Exception as e:
        print(f"Lỗi xử lý: {e}")
        # Tạo document lỗi
        error_doc = Document()
        error_doc.add_heading('LỖI XỬ LÝ', level=1)
        error_doc.add_paragraph(f"Có lỗi xảy ra: {str(e)}")
        error_doc.add_paragraph("Vui lòng kiểm tra lại file đầu vào.")
        
        stats = {
            'latex_count': 0,
            'table_count': 0,
            'processed_paragraphs': 0,
            'error': str(e)
        }
        
        return error_doc, stats
