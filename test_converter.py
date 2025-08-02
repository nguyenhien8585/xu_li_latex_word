#!/usr/bin/env python3
"""
Test script cho Enhanced Word Converter
Tạo file Word mẫu với các công thức LaTeX để test
"""

import os
from docx import Document
from datetime import datetime

def create_test_document():
    """Tạo document Word mẫu với các công thức test"""
    
    doc = Document()
    
    # Title
    title = doc.add_heading('Test Document - LaTeX to Word Equations', 0)
    
    # Thêm metadata
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("Purpose: Test Enhanced Word Converter equation capabilities")
    
    doc.add_heading('1. Basic Math Formulas', level=1)
    
    # Test cases cơ bản
    test_cases = [
        ("Einstein's Formula", "$E = mc^{2}$"),
        ("Pythagorean Theorem", "$a^{2} + b^{2} = c^{2}$"),
        ("Quadratic Formula", "$x = \\frac{-b \\pm \\sqrt{b^{2} - 4ac}}{2a}$"),
        ("Simple Fraction", "$\\frac{1}{2} + \\frac{1}{3} = \\frac{5}{6}$"),
        ("Square Root", "$\\sqrt{x^{2} + y^{2}} = r$"),
        ("Subscript", "$x_{1} + x_{2} + ... + x_{n}$"),
        ("Greek Letters", "$\\alpha + \\beta = \\gamma$"),
        ("Trigonometry", "$\\sin^{2}\\theta + \\cos^{2}\\theta = 1$"),
        ("Calculus", "$\\frac{d}{dx}(x^{n}) = nx^{n-1}$"),
        ("Limit", "$\\lim_{x \\to \\infty} \\frac{1}{x} = 0$"),
    ]
    
    for description, formula in test_cases:
        p = doc.add_paragraph()
        p.add_run(f"{description}: ")
        p.add_run(formula)
    
    doc.add_heading('2. Complex Formulas', level=1)
    
    complex_cases = [
        ("Integral", "$\\int_{a}^{b} f(x)dx = F(b) - F(a)$"),
        ("Summation", "$\\sum_{i=1}^{n} i = \\frac{n(n+1)}{2}$"),
        ("Matrix", "$A = \\begin{pmatrix} a & b \\\\ c & d \\end{pmatrix}$"),
        ("Probability", "$P(A|B) = \\frac{P(B|A)P(A)}{P(B)}$"),
        ("Complex Fraction", "$\\frac{\\frac{a}{b}}{\\frac{c}{d}} = \\frac{ad}{bc}$"),
        ("Nested Root", "$\\sqrt{\\sqrt{x} + \\sqrt{y}}$"),
        ("Multiple Subscripts", "$x_{i,j}^{(k)}$"),
        ("Physics Formula", "$F = ma = m\\frac{dv}{dt}$"),
    ]
    
    for description, formula in complex_cases:
        p = doc.add_paragraph()
        p.add_run(f"{description}: ")
        p.add_run(formula)
    
    doc.add_heading('3. Inline Math with Text', level=1)
    
    inline_examples = [
        "The famous formula $E = mc^{2}$ was derived by Einstein.",
        "In a right triangle, we have $a^{2} + b^{2} = c^{2}$ where $c$ is the hypotenuse.",
        "The derivative of $x^{n}$ is $nx^{n-1}$, which is a fundamental rule in calculus.",
        "For any real numbers $a$ and $b$, we have $(a+b)^{2} = a^{2} + 2ab + b^{2}$.",
        "The area of a circle is $A = \\pi r^{2}$ where $r$ is the radius.",
    ]
    
    for example in inline_examples:
        doc.add_paragraph(example)
    
    doc.add_heading('4. Markdown Table Test', level=1)
    
    table_text = """
| Formula Name | LaTeX | Description |
|--------------|-------|-------------|
| Quadratic | $ax^{2} + bx + c = 0$ | Second degree polynomial |
| Exponential | $e^{x}$ | Natural exponential function |
| Logarithm | $\\ln(x)$ | Natural logarithm |
| Trigonometric | $\\sin(\\theta)$ | Sine function |
"""
    
    doc.add_paragraph(table_text)
    
    doc.add_heading('5. Special Characters & Symbols', level=1)
    
    symbols_cases = [
        ("Infinity", "$\\lim_{x \\to \\infty} f(x) = \\infty$"),
        ("Plus/Minus", "$x = a \\pm b$"),
        ("Not Equal", "$a \\neq b$"),
        ("Less/Greater Equal", "$a \\leq b \\leq c \\geq d$"),
        ("Approximately", "$\\pi \\approx 3.14159$"),
        ("Multiplication", "$a \\times b = ab$"),
        ("Division", "$a \\div b = \\frac{a}{b}$"),
        ("Partial Derivative", "$\\frac{\\partial f}{\\partial x}$"),
        ("Nabla", "$\\nabla f = (\\frac{\\partial f}{\\partial x}, \\frac{\\partial f}{\\partial y})$"),
    ]
    
    for description, formula in symbols_cases:
        p = doc.add_paragraph()
        p.add_run(f"{description}: ")
        p.add_run(formula)
    
    doc.add_heading('6. Edge Cases', level=1)
    
    edge_cases = [
        ("Empty Formula", "${}$"),
        ("Single Character", "$x$"),
        ("Numbers Only", "$123$"),
        ("Mixed Brackets", "$\\left(\\frac{a}{b}\\right)^{2}$"),
        ("Multiple Formulas", "$a = b$ and $c = d$ in same line"),
        ("Formula with Text", "$f(x) = ax + b$ where $a, b \\in \\mathbb{R}$"),
    ]
    
    for description, formula in edge_cases:
        p = doc.add_paragraph()
        p.add_run(f"{description}: ")
        p.add_run(formula)
    
    # Instructions
    doc.add_page_break()
    doc.add_heading('Testing Instructions', level=1)
    
    instructions = [
        "1. Save this document as 'test_input.docx'",
        "2. Upload to Enhanced Word Converter app",
        "3. Click 'Chuyển đổi thành Equation Word'",
        "4. Download the converted file",
        "5. Open in Microsoft Word",
        "6. Verify that formulas are real equations (can double-click to edit)",
        "7. Check that inline text is preserved",
        "8. Verify table conversion worked",
        "9. Test equation editing functionality",
        "10. Compare with original to ensure accuracy"
    ]
    
    for instruction in instructions:
        doc.add_paragraph(instruction)
    
    doc.add_heading('Expected Results', level=1)
    
    results = [
        "✅ All $...$ formulas converted to Word equations",
        "✅ Equations are editable (double-click works)",
        "✅ Fractions display as proper fractions",
        "✅ Square roots display with radical symbol",
        "✅ Superscripts and subscripts formatted correctly",
        "✅ Greek letters display as symbols",
        "✅ Inline text preserved around equations",
        "✅ Markdown table converted to Word table",
        "✅ Special symbols display correctly",
        "✅ Complex formulas maintain structure"
    ]
    
    for result in results:
        doc.add_paragraph(result)
    
    return doc

def main():
    """Main function để tạo test document"""
    
    print("🧪 Creating test document for Enhanced Word Converter...")
    
    # Tạo document
    doc = create_test_document()
    
    # Save file
    filename = "test_input_enhanced.docx"
    doc.save(filename)
    
    print(f"✅ Test document created: {filename}")
    print(f"📁 File size: {os.path.getsize(filename) / 1024:.1f} KB")
    
    # Thống kê
    formulas_count = 0
    for para in doc.paragraphs:
        formulas_count += para.text.count('$')
    
    print(f"📊 Statistics:")
    print(f"   - Total paragraphs: {len(doc.paragraphs)}")
    print(f"   - Formula markers ($): {formulas_count}")
    print(f"   - Estimated formulas: {formulas_count // 2}")
    
    print("\n🚀 Next steps:")
    print("1. Upload this file to Enhanced Word Converter")
    print("2. Process and download result")
    print("3. Open in Word and verify equations")
    print("4. Check editability of equations")
    
    return filename

if __name__ == "__main__":
    filename = main()
    print(f"\n📄 Test file ready: {filename}")
